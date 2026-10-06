"""Assembles the final report bundle and renders Markdown / PDF / DOCX exports."""
from __future__ import annotations

import io
import re
from html import escape

from ..db import Database
from .citation_service import render_markers

TIER_NAMES = {1: "Tier 1 (peer-reviewed / official)", 2: "Tier 2 (academic repository)",
              3: "Tier 3 (industry / quality press)", 4: "Tier 4 (blog / forum)"}


def final_bundle(db: Database, sid: str) -> dict | None:
    sess = db.get_session(sid)
    reports = db.get_reports(sid)
    if not sess or not reports:
        return None
    crit = db.get_stage_output(sid, "critique")
    version = crit["final_version"] if crit else reports[-1]["version"]
    report = next((r for r in reports if r["version"] == version), reports[-1])
    return {"session": sess, "report": report, "critique": crit,
            "sources": db.get_items("sources", sid), "plan": db.get_stage_output(sid, "planning"),
            "is_final": bool(crit)}


def pretty_date(iso: str | None) -> str:
    from datetime import datetime
    try:
        return datetime.fromisoformat(iso or "").strftime("%d %b %Y, %H:%M UTC")
    except ValueError:
        return iso or ""


def render_text(report: dict, style: str, text: str) -> str:
    return render_markers(text, report.get("labels", {}), style)


def to_markdown(b: dict) -> str:
    s, r = b["session"], b["report"]
    style = s["citation_style"]
    out = []
    if s["mode"] == "demo":
        out.append("> **DEMO MODE — SIMULATED AI. Sample data, not real sources or real AI research.**\n")
    out.append(f"# {r['title']}\n")
    out.append(f"*Report type:* {s['report_type']}  \n*Citation style:* {style}  \n"
               f"*Generated:* {pretty_date(s['completed_at'] or s['updated_at'])}  \n"
               f"*Quality score:* {s['quality_score'] if s['quality_score'] is not None else 'n/a'}/100\n")
    for sec in r["sections"]:
        out.append(f"## {sec['heading']}\n\n{render_text(r, style, sec['content'])}\n")
    out.append("## References\n")
    for ref in r["references"]:
        out.append(f"{ref['text']}\n" if style == "IEEE" else f"- {ref['text']}")
    used = set(r.get("citation_map", {}))
    out.append("\n## Appendix: Source Credibility\n")
    out.append("| ID | Title | Tier | Type | Retrieval |\n|---|---|---|---|---|")
    for src in b["sources"]:
        if src["source_id"] in used:
            out.append(f"| {src['source_id']} | {src['title'][:80].replace('|', '/')} | "
                       f"{src['credibility_tier']} | {src['source_type']} | {src['fetch_status']} |")
    return "\n".join(out) + "\n"


def _inline_html(text: str) -> str:
    t = escape(text)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", t)
    return t


def _blocks(content: str):
    """Yield ('p', text) or ('li', text) from markdown-lite content."""
    for para in re.split(r"\n\s*\n", content.strip()):
        lines = [l for l in para.split("\n") if l.strip()]
        if lines and all(re.match(r"^\s*[-*]\s+", l) for l in lines):
            for l in lines:
                yield "li", re.sub(r"^\s*[-*]\s+", "", l)
        else:
            yield "p", " ".join(l.strip() for l in lines)


def to_pdf(b: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    s, r = b["session"], b["report"]
    style = s["citation_style"]
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="Times-Roman", fontSize=11,
                          leading=15, alignment=TA_JUSTIFY, spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName="Times-Bold", fontSize=18, leading=22)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Times-Bold", fontSize=13,
                        spaceBefore=12, spaceAfter=4, textColor=colors.HexColor("#1f3a5f"))
    meta = ParagraphStyle("meta", parent=body, fontSize=9, leading=12, textColor=colors.HexColor("#555555"),
                          alignment=0)
    ref = ParagraphStyle("ref", parent=body, fontSize=9.5, leading=12.5, alignment=0,
                         leftIndent=18, firstLineIndent=-18)
    small = ParagraphStyle("small", parent=meta, fontSize=8, leading=10)

    def page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Times-Roman", 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        footer = "Multi-Agent Research & Report Generator"
        if s["mode"] == "demo":
            footer = "DEMO MODE — SIMULATED AI (sample data). " + footer
        canvas.drawString(2 * cm, 1.2 * cm, footer)
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    flow = []
    if s["mode"] == "demo":
        flow.append(Paragraph("<b>DEMO MODE — SIMULATED AI. Sample data, not real sources or real AI research.</b>",
                              ParagraphStyle("demo", parent=body, textColor=colors.HexColor("#a33"))))
    flow += [Paragraph(escape(r["title"]), h1),
             Paragraph(f"{escape(s['report_type'])} &nbsp;|&nbsp; Citation style: {style} &nbsp;|&nbsp; "
                       f"Quality score: {s['quality_score'] if s['quality_score'] is not None else 'n/a'}/100"
                       f" &nbsp;|&nbsp; {escape(pretty_date(s['completed_at'] or s['updated_at']))}", meta),
             Spacer(1, 10)]
    for sec in r["sections"]:
        flow.append(Paragraph(escape(sec["heading"]), h2))
        items = []
        for kind, txt in _blocks(render_text(r, style, sec["content"])):
            if kind == "li":
                items.append(ListItem(Paragraph(_inline_html(txt), body), leftIndent=12))
                continue
            if items:
                flow.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
                items = []
            flow.append(Paragraph(_inline_html(txt), body))
        if items:
            flow.append(ListFlowable(items, bulletType="bullet", start="•", leftIndent=12))
    flow.append(Paragraph("References", h2))
    for rf in r["references"]:
        flow.append(Paragraph(_inline_html(rf["text"]), ref))
    used = set(r.get("citation_map", {}))
    rows = [["ID", "Title", "Tier", "Type", "Retrieval"]]
    for src in b["sources"]:
        if src["source_id"] in used:
            rows.append([src["source_id"], Paragraph(escape(src["title"][:110]), small),
                         str(src["credibility_tier"]), src["source_type"], src["fetch_status"]])
    if len(rows) > 1:
        flow += [Paragraph("Appendix: Source Credibility", h2)]
        t = Table(rows, colWidths=[1.2 * cm, 9.3 * cm, 1.1 * cm, 2.9 * cm, 2.3 * cm], repeatRows=1)
        t.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, 0), "Times-Bold", 9), ("FONT", (0, 1), (-1, -1), "Times-Roman", 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        flow.append(t)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title=r["title"])
    doc.build(flow, onFirstPage=page, onLaterPages=page)
    return buf.getvalue()


def _docx_runs(par, text: str) -> None:
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            par.add_run(chunk).bold = i % 2 == 1


def to_docx(b: dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    s, r = b["session"], b["report"]
    style = s["citation_style"]
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    if s["mode"] == "demo":
        p = doc.add_paragraph()
        run = p.add_run("DEMO MODE — SIMULATED AI. Sample data, not real sources or real AI research.")
        run.bold = True
        run.font.color.rgb = RGBColor(0xA3, 0x33, 0x33)
    doc.add_heading(r["title"], level=0)
    m = doc.add_paragraph(f"{s['report_type']} | Citation style: {style} | Quality score: "
                          f"{s['quality_score'] if s['quality_score'] is not None else 'n/a'}/100 | "
                          f"{pretty_date(s['completed_at'] or s['updated_at'])}")
    m.runs[0].font.size = Pt(9)
    for sec in r["sections"]:
        doc.add_heading(sec["heading"], level=1)
        for kind, txt in _blocks(render_text(r, style, sec["content"])):
            p = doc.add_paragraph(style="List Bullet" if kind == "li" else None)
            if kind == "p":
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _docx_runs(p, txt)
    doc.add_heading("References", level=1)
    for rf in r["references"]:
        p = doc.add_paragraph(rf["text"])
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        for run in p.runs:
            run.font.size = Pt(10)
    used = set(r.get("citation_map", {}))
    srcs = [x for x in b["sources"] if x["source_id"] in used]
    if srcs:
        doc.add_heading("Appendix: Source Credibility", level=1)
        t = doc.add_table(rows=1, cols=5)
        t.style = "Light Grid Accent 1"
        for i, h in enumerate(["ID", "Title", "Tier", "Type", "Retrieval"]):
            t.rows[0].cells[i].text = h
        for x in srcs:
            cells = t.add_row().cells
            for i, v in enumerate([x["source_id"], x["title"][:110], str(x["credibility_tier"]),
                                   x["source_type"], x["fetch_status"]]):
                cells[i].text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:60] or "report"
    return f"{base}.{ext}"
